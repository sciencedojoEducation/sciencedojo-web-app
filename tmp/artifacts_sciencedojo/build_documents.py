from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/Users/piumalmahawasala/Documents/Codex/sciencedojo-web-app')
OUT = ROOT / 'deliverables/sciencedojo_triathlon_pack'
LOGO = ROOT / 'public/images/sciencedojo-wordmark-gradient.jpg'

NAVY = '11243E'
BLUE = '3973E6'
CYAN = '2FC4C9'
INK = '1F2937'
MUTED = '607086'
PALE = 'EEF4FF'
PALE_CYAN = 'EAFBFB'
WHITE = 'FFFFFF'
LINE = 'CAD7E8'


def font(run, size=10.5, bold=False, color=INK, name='Aptos'):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for edge, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{edge}'))
        if node is None:
            node = OxmlElement(f'w:{edge}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_cell_border(cell, color=LINE, size='8'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = qn(f'w:{edge}')
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f'w:{edge}')
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), size)
        el.set(qn('w:color'), color)


def set_table_geometry(table, widths_inches, indent_dxa=120):
    table.autofit = False
    total = sum(widths_inches)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in('w:tblW')
    tbl_w.set(qn('w:w'), str(round(total * 1440)))
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_ind = tbl_pr.first_child_found_in('w:tblInd')
    if tbl_ind is None:
        tbl_ind = OxmlElement('w:tblInd')
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn('w:w'), str(indent_dxa))
    tbl_ind.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_inches:
        grid_col = OxmlElement('w:gridCol')
        grid_col.set(qn('w:w'), str(round(width * 1440)))
        grid.append(grid_col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx >= len(widths_inches):
                continue
            dxa = round(widths_inches[idx] * 1440)
            cell.width = Inches(widths_inches[idx])
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in('w:tcW')
            tc_w.set(qn('w:w'), str(dxa))
            tc_w.set(qn('w:type'), 'dxa')


def set_page(doc, landscape=False, margins=(0.62, 0.62, 0.58, 0.62)):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    if landscape:
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
    else:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
    top, right, bottom, left = margins
    section.top_margin = Inches(top)
    section.right_margin = Inches(right)
    section.bottom_margin = Inches(bottom)
    section.left_margin = Inches(left)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    return section


def base_styles(doc, compact=False):
    normal = doc.styles['Normal']
    normal.font.name = 'Aptos'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    normal.font.size = Pt(9.4 if compact else 10.3)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(4 if compact else 5)
    normal.paragraph_format.line_spacing = 1.08 if compact else 1.12
    for style_name, size, color, before, after in [
        ('Title', 25, NAVY, 0, 4),
        ('Subtitle', 11.5, MUTED, 0, 10),
        ('Heading 1', 14, BLUE, 9, 3),
        ('Heading 2', 11.5, NAVY, 6, 2),
        ('Heading 3', 10.5, NAVY, 4, 2),
    ]:
        style = doc.styles[style_name]
        style.font.name = 'Aptos Display' if style_name in ('Title', 'Heading 1') else 'Aptos'
        style._element.rPr.rFonts.set(qn('w:ascii'), style.font.name)
        style._element.rPr.rFonts.set(qn('w:hAnsi'), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = style_name != 'Subtitle'
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_logo(doc, width=0.78):
    if LOGO.exists():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        picture = p.add_run().add_picture(str(LOGO), width=Inches(width))
        picture._inline.docPr.set('descr', 'ScienceDojo wordmark')


def add_label_para(container, label, text, size=9.5, after=3):
    p = container.add_paragraph() if hasattr(container, 'add_paragraph') else container
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(label + ' ')
    font(r, size=size, bold=True, color=NAVY)
    r = p.add_run(text)
    font(r, size=size, color=INK)
    return p


def add_footer(section, text):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(text)
    font(r, size=8, color=MUTED)


def build_onepager():
    doc = Document()
    section = set_page(doc, landscape=False)
    base_styles(doc)
    add_logo(doc)

    title = doc.add_paragraph(style='Title')
    title.add_run('ScienceDojo')
    sub = doc.add_paragraph(style='Subtitle')
    sub.add_run('Human-first STEM tutoring, structured from first assessment to confident independent learning')

    lead = doc.add_paragraph()
    lead.paragraph_format.space_after = Pt(7)
    lead.paragraph_format.line_spacing = 1.15
    r = lead.add_run('The opportunity. ')
    font(r, size=11, bold=True, color=BLUE)
    r = lead.add_run('Many capable students understand more than their grades show. Families often receive isolated lessons, fragmented feedback, and little visibility into what should happen between sessions.')
    font(r, size=11, color=NAVY)

    doc.add_heading('What ScienceDojo delivers', level=1)
    add_label_para(doc, 'Verified tutoring.', 'Families discover and book STEM tutors through a supported platform with onboarding, credential review, safeguarding processes, payments, and dispute handling.')
    add_label_para(doc, 'Continuity between lessons.', 'Class spaces connect lesson notes, homework, messaging, progress records, and live learning tools so support does not end when the call ends.')
    add_label_para(doc, 'Guided practice.', 'Practice Dojo offers open curriculum-aligned knowledge checks; premium Missions turn lesson context into personalised recall, reasoning, application, and error-correction tasks.')
    add_label_para(doc, 'Parent clarity.', 'Parents can see what was covered, what comes next, and how confidence and consistency are developing.')

    doc.add_heading('Who it serves', level=1)
    audience = doc.add_paragraph()
    audience.paragraph_format.space_after = Pt(5)
    for label, text in [
        ('Students', ' preparing for GCSE, IGCSE, A-Level and related STEM pathways.'),
        ('Parents', ' seeking trusted tutors, reliable communication, and visible learning progress.'),
        ('Tutors', ' who want qualified leads and an operating system for consistent teaching support.'),
    ]:
        r = audience.add_run(label)
        font(r, size=9.7, bold=True, color=NAVY)
        r = audience.add_run(text + '  ')
        font(r, size=9.7)

    doc.add_heading('Business model', level=1)
    add_label_para(doc, 'Core revenue.', 'A platform commission on completed tutoring bookings; the current platform setting defaults to 25%.')
    add_label_para(doc, 'Additional revenue.', 'FocusDojo Pro at €4.99/month or €39/year, with future potential for school partnerships and premium learning programmes.')

    doc.add_heading('Why it can win', level=1)
    add_label_para(doc, 'A learning system, not only a marketplace.', 'The compounding asset is the connected record of assessment, lesson activity, tutor feedback, homework, practice, and progress.')
    add_label_para(doc, 'Human expertise stays central.', 'AI supports question generation and structured follow-up, while tutors remain the trusted relationship and learning authority.')
    add_label_para(doc, 'Operational foundations already exist.', 'The product includes multi-role dashboards, booking, payments and payouts, class spaces, messaging, reviews, tutor onboarding, safeguarding workflows, and administrative controls.')

    doc.add_heading('Near-term validation focus', level=1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    for idx, text in enumerate([
        'Convert free assessments into first paid lessons.',
        'Increase repeat bookings and mission completion between lessons.',
        'Prove parent-reported confidence and tutor retention.',
        'Pilot a university or regional partnership channel.',
    ], 1):
        r = p.add_run(f'{idx}. ')
        font(r, size=9.4, bold=True, color=CYAN)
        r = p.add_run(text + ('   ' if idx < 4 else ''))
        font(r, size=9.4, color=INK)

    add_footer(section, 'ScienceDojo | General company overview | September 2026')
    path = OUT / 'ScienceDojo_OnePager.docx'
    doc.save(path)
    return path


def clear_cell(cell):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    return p


def add_canvas_block(cell, title, items, fill=WHITE):
    shade(cell, fill)
    set_cell_border(cell)
    set_cell_margins(cell, top=105, start=125, bottom=95, end=125)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    p = clear_cell(cell)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    font(r, size=11.5, bold=True, color=BLUE)
    for label, text in items:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2.2)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(label + ': ')
        font(r, size=8.2, bold=True, color=NAVY)
        r = p.add_run(text)
        font(r, size=8.2, color=INK)


def build_canvas():
    doc = Document()
    section = set_page(doc, landscape=True, margins=(0.38, 0.40, 0.36, 0.40))
    base_styles(doc, compact=True)

    header = doc.add_table(rows=1, cols=2)
    header.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(header, [2.0, 8.0], indent_dxa=0)
    set_repeat_table_header(header.rows[0])
    for c in header.rows[0].cells:
        set_cell_margins(c, top=0, start=0, bottom=0, end=0)
        tc_pr = c._tc.get_or_add_tcPr()
        borders = OxmlElement('w:tcBorders')
        for edge in ('top', 'left', 'bottom', 'right'):
            el = OxmlElement(f'w:{edge}')
            el.set(qn('w:val'), 'nil')
            borders.append(el)
        tc_pr.append(borders)
    if LOGO.exists():
        p = header.cell(0, 0).paragraphs[0]
        picture = p.add_run().add_picture(str(LOGO), width=Inches(0.72))
        picture._inline.docPr.set('descr', 'ScienceDojo wordmark')
    p = header.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run('BUSINESS MODEL CANVAS')
    font(r, size=21, bold=True, color=NAVY, name='Aptos Display')
    p2 = header.cell(0, 1).add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p2.paragraph_format.space_after = Pt(4)
    r = p2.add_run('Working model - assumptions to validate through pilots and customer evidence')
    font(r, size=8.5, color=MUTED)

    table = doc.add_table(rows=2, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    widths = [2.0, 2.0, 2.0, 2.0, 2.0]
    set_table_geometry(table, widths, indent_dxa=0)
    set_repeat_table_header(table.rows[0])

    add_canvas_block(table.cell(0, 0), 'Key Partners', [
        ('Supply', 'Verified STEM tutors and specialist educators'),
        ('Infrastructure', 'Supabase, Stripe, Resend, video classroom and AI providers'),
        ('Channels', 'Schools, universities, regional programmes and parent communities'),
        ('Trust', 'Credential verification and safeguarding partners'),
    ], PALE)
    add_canvas_block(table.cell(0, 1), 'Key Activities', [
        ('Acquire', 'Assess learners, recruit tutors and build partner referrals'),
        ('Deliver', 'Match, book, teach, record, assign and review'),
        ('Assure', 'Tutor verification, safeguarding, payments and dispute operations'),
        ('Improve', 'Use learning records to strengthen personalisation and tutor quality'),
    ])
    add_canvas_block(table.cell(0, 2), 'Value Propositions', [
        ('Students', 'Clear explanations, guided practice and confidence that compounds'),
        ('Parents', 'Trusted tutors, organised support and visibility into the next step'),
        ('Tutors', 'Qualified demand plus tools for consistent, accountable teaching'),
        ('Partners', 'A measurable human-first STEM support pathway'),
    ], PALE_CYAN)
    add_canvas_block(table.cell(0, 3), 'Customer Relationships', [
        ('Entry', 'Free learning assessment and open Practice Dojo'),
        ('Core', 'Ongoing one-to-one tutor relationship'),
        ('Continuity', 'Class space, messaging, homework, Missions and progress updates'),
        ('Retention', 'Recurring lessons, parent reassurance and demonstrated momentum'),
    ])
    add_canvas_block(table.cell(0, 4), 'Customer Segments', [
        ('Primary buyer', 'Parents of secondary-school students needing STEM support'),
        ('Primary user', 'GCSE, IGCSE and A-Level learners'),
        ('Supply user', 'Independent and university-linked STEM tutors'),
        ('Future buyer', 'Schools, universities and regional education programmes'),
    ], PALE)

    add_canvas_block(table.cell(1, 0), 'Key Resources', [
        ('People', 'Tutor network, educational operations and safeguarding capability'),
        ('Product', 'Multi-role platform, class spaces, booking, payments and learning tools'),
        ('Data', 'Consent-based lesson, practice, homework and progress records'),
        ('Brand', 'Human-first trust and structured science-confidence positioning'),
    ])
    add_canvas_block(table.cell(1, 1), 'Channels', [
        ('Direct', 'SEO learning content, tutor pages, assessment funnel and referrals'),
        ('Product-led', 'Practice Dojo and FocusDojo'),
        ('Partner-led', 'University, school and regional programme pilots'),
        ('Community', 'Exam-support community and tutor advocacy'),
    ], PALE)
    add_canvas_block(table.cell(1, 2), 'Cost Structure', [
        ('Variable', 'Tutor payouts, payment fees, AI usage, email and video services'),
        ('Operating', 'Verification, safeguarding, support and dispute handling'),
        ('Growth', 'Tutor recruitment, assessments, partnerships and content'),
        ('Product', 'Engineering, hosting, security and compliance'),
    ])
    add_canvas_block(table.cell(1, 3), 'Revenue Streams', [
        ('Tutoring', 'Platform commission on completed bookings; current default is 25%'),
        ('Subscription', 'FocusDojo Pro: €4.99 monthly or €39 yearly'),
        ('Potential', 'School or programme licences, cohort support and premium pathways'),
        ('Principle', 'Keep tutoring as the core; add software revenue without weakening trust'),
    ], PALE_CYAN)
    add_canvas_block(table.cell(1, 4), 'Critical Metrics', [
        ('Funnel', 'Assessment-to-first-paid-lesson conversion'),
        ('Retention', 'Repeat-booking rate and active learner months'),
        ('Learning', 'Homework/Mission completion and confidence improvement'),
        ('Supply', 'Tutor activation, utilisation, quality and retention'),
    ])

    add_footer(section, 'ScienceDojo | Business Model Canvas | September 2026')
    path = OUT / 'ScienceDojo_Business_Model_Canvas.docx'
    doc.save(path)
    return path


if __name__ == '__main__':
    OUT.mkdir(parents=True, exist_ok=True)
    for p in [build_onepager(), build_canvas()]:
        print(p)
